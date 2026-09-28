import io
import logging
import pdfplumber
import docx
from fastapi import HTTPException
from app.core.security import sanitize_untrusted_input
from app.core.llm import llm_client
from app.models.schemas import ResumeAnalysisResult
from app.prompts.resume_prompt import RESUME_ANALYSIS_SYSTEM_PROMPT, RESUME_ANALYSIS_USER_PROMPT

logger = logging.getLogger("uvicorn.error")

def extract_text_from_pdf(content: bytes) -> str:
    """Extract plain text from PDF file bytes using pdfplumber."""
    extracted_pages = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                extracted_pages.append(text)
    return "\n".join(extracted_pages).strip()

def extract_text_from_docx(content: bytes) -> str:
    """Extract plain text from DOCX file bytes using python-docx."""
    doc = docx.Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text from table cells if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
                
    return "\n".join(paragraphs).strip()

async def parse_resume_bytes(filename: str, content: bytes) -> ResumeAnalysisResult:
    """Extract text from resume PDF/DOCX and run structured LLM analysis."""
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    
    try:
        if ext == "pdf":
            raw_text = extract_text_from_pdf(content)
        elif ext == "docx":
            raw_text = extract_text_from_docx(content)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Must be PDF or DOCX.")
    except Exception as e:
        logger.error(f"Error parsing resume file '{filename}': {e}")
        raise HTTPException(status_code=400, detail=f"Failed to read resume document: {str(e)}")

    if not raw_text or len(raw_text.strip()) < 20:
        raise HTTPException(status_code=400, detail="Uploaded resume document appears to be empty or unreadable text.")

    sanitized = sanitize_untrusted_input(raw_text)
    prompt = RESUME_ANALYSIS_USER_PROMPT.format(sanitized_resume_text=sanitized)

    result = await llm_client.generate_json(
        prompt=prompt,
        schema_cls=ResumeAnalysisResult,
        system_prompt=RESUME_ANALYSIS_SYSTEM_PROMPT
    )
    return result
