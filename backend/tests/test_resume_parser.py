import pytest
from app.services.resume_parser import parse_resume_bytes

@pytest.mark.asyncio
async def test_parse_docx_resume():
    # Construct a minimal DOCX in memory
    import docx, io
    doc = docx.Document()
    doc.add_heading('John Doe', 0)
    doc.add_paragraph('Software Engineer with 5 years experience in Python, FastAPI, and React.')
    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()

    result = await parse_resume_bytes("sample_resume.docx", content)
    assert result is not None
    assert result.experience_years >= 0.0
    assert isinstance(result.skills, list)
