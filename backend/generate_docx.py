import os
import glob
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_master_docx():
    doc = Document()

    # Document Title Styling
    title = doc.add_heading(level=0)
    run = title.add_run("HireFit AI: Adaptive Target-Company & Role-Specific Interview Simulator")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

    sub = doc.add_paragraph()
    sub_run = sub.add_run("Comprehensive Technical Architecture, System Specifications & Design Document")
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("Version: 1.0.0 | Production Release | Author: Senior Technical Writer & Lead Systems Architect\n")
    doc.add_page_break()

    # Process all markdown files in order
    docs_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
    md_files = sorted(glob.glob(os.path.join(docs_dir, "*.md")))

    for md_path in md_files:
        filename = os.path.basename(md_path)
        print(f"Converting {filename} to DOCX...")

        with open(md_path, "r", encoding="utf-8") as f:
            lines = f.readlines()

        in_code_block = False
        code_lines = []

        for line in lines:
            line_str = line.rstrip()

            # Handle Code Blocks
            if line_str.startswith("```"):
                if in_code_block:
                    # End code block
                    code_text = "\n".join(code_lines)
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.left_indent = Inches(0.3)
                    run = p.add_run(code_text)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(30, 41, 59)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lines = []
                continue

            if in_code_block:
                code_lines.append(line_str)
                continue

            # Headings
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str.startswith("#### "):
                doc.add_heading(line_str[5:], level=4)

            # Bullet points
            elif line_str.startswith("- ") or line_str.startswith("* "):
                doc.add_paragraph(line_str[2:], style='List Bullet')
            elif re.match(r'^\d+\.\s+', line_str):
                text_part = re.sub(r'^\d+\.\s+', '', line_str)
                doc.add_paragraph(text_part, style='List Number')

            # Horizontal Rule
            elif line_str == "---":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(12)

            # Plain paragraph text
            elif line_str.strip():
                # Strip simple bold markdown tags for clean Word rendering
                clean_text = line_str.replace("**", "").replace("*", "")
                p = doc.add_paragraph(clean_text)
                p.paragraph_format.space_after = Pt(4)

        doc.add_page_break()

    out_path = os.path.join(docs_dir, "HireFit_AI_Comprehensive_Project_Documentation.docx")
    doc.save(out_path)
    print(f"\n🎉 Successfully generated DOCX documentation at: {out_path}")
    return out_path

if __name__ == "__main__":
    create_master_docx()
